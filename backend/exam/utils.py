import uuid
import docx
import re
import openpyxl

def parse_questions(filepath, filename, uid=None):
    if uid is None:
        uid = str(uuid.uuid4())
    questions = []
    total_score = 0
    content = ""
    if filename.endswith('.txt'):
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
    elif filename.endswith('.docx'):
        doc = docx.Document(filepath)
        content = "\n".join([p.text for p in doc.paragraphs])
    
    lines = content.split('\n')
    current_question = None
    current_options = []
    correct_answer = None
    question_type = None  # 'mcq' or 'coding'
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        if line.startswith("Problem Statement:"):
            if current_question:
                score = 2 if question_type == 'mcq' else 5
                total_score += score
                questions.append({
                    "type": question_type,
                    "question": current_question,
                    "options": current_options if question_type == 'mcq' else [],
                    "correctAnswer": correct_answer if question_type == 'mcq' else None,
                    "score": score
                })
                current_question, current_options, correct_answer = None, [], None
            question_type = 'coding'
            current_question = line
            continue

        if line.endswith('?'):
            if current_question:
                score = 2 if question_type == 'mcq' else 5
                total_score += score
                questions.append({
                    "type": question_type,
                    "question": current_question,
                    "options": current_options if question_type == 'mcq' else [],
                    "correctAnswer": correct_answer if question_type == 'mcq' else None,
                    "score": score
                })
                current_question, current_options, correct_answer = None, [], None
            question_type = 'mcq'
            current_question = line
            continue

        if question_type == 'mcq' and (line.startswith("a)") or line.startswith("b)") or 
            line.startswith("c)") or line.startswith("d)")):
            current_options.append(line)
            continue

        if line.lower().startswith("correct:"):
            match = re.search(r'\((.*?)\)', line)
            if match:
                correct_answer = match.group(1).strip()
            continue

        if current_question:
            current_question += " " + line
        else:
            current_question = line

    if current_question:
        score = 2 if question_type == 'mcq' else 5
        total_score += score
        questions.append({
            "type": question_type,
            "question": current_question,
            "options": current_options if question_type == 'mcq' else [],
            "correctAnswer": correct_answer if question_type == 'mcq' else None,
            "score": score
        })

    return {"examId": uid, "questions": questions, "maxScore": total_score}


def parse_student_excel(filepath):
    """Parses an Excel file containing student data.
       Assumes the first row contains headers, e.g. Name, Username, Email.
    """
    wb = openpyxl.load_workbook(filepath)
    sheet = wb.active
    students = []
    # Skip header row (assumed row 1)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] and row[1] and row[2]:
            student = {"name": row[0], "username": row[1], "email": row[2]}
            students.append(student)
    return students